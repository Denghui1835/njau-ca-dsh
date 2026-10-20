"""校验我们写出的 .docx 能被独立解析器（python-docx）正常打开。

这是「不是自说自话」的关键一步：我们自己写、自己读当然能通过；
换一个完全不共享代码的实现来读，才能证明产出是标准 OOXML 文档。
"""

import sys
from docx import Document


def dump(path: str) -> None:
    doc = Document(path)
    paras = [p for p in doc.paragraphs]
    non_empty = [p for p in paras if p.text.strip()]
    tables = doc.tables
    print(f"  file        : {path}")
    print(f"  paragraphs  : {len(paras)} (有字 {len(non_empty)})")
    print(f"  tables      : {len(tables)}")
    if tables:
        first = tables[0]
        print(f"  table[0]    : {len(first.rows)} 行 x {len(first.columns)} 列")
        cell = first.rows[0].cells[0].text.strip()
        print(f"  单元格[0,0] : {cell[:40]!r}")
    for p in non_empty[:5]:
        style = p.style.name if p.style is not None else "?"
        print(f"  段[{style:>10}] {p.text[:56]!r}")
    print()


def main() -> int:
    if len(sys.argv) < 2:
        print("用法: python verify_docx.py <file.docx> [...]")
        return 2
    failed = 0
    for path in sys.argv[1:]:
        try:
            dump(path)
        except Exception as error:  # noqa: BLE001 - 这里就是要抓一切
            failed += 1
            print(f"  ✗ {path} 打不开: {error}")
    print("全部可被 python-docx 打开" if failed == 0 else f"{failed} 个文件有问题")
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
